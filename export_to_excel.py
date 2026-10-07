"""
export_to_excel.py
Step 10 - combines everything the pipeline has built into one Excel
workbook: detected_possessions_final.json (corrected boundaries) +
possession_results.json (tagged outcomes) + formation_analysis.json +
cuts_analysis.json, joined by (quarter, team, start_sec) - start_sec is
used as the join key rather than the full possession ID, since a tagging
correction changes end_sec but never start_sec.

Built as a SCOUTING REPORT for coaches, not analysts: sheets are ordered
simple-to-complex (overview first, full per-possession detail last), and
the two teams are NEVER combined/blended anywhere - every sheet keeps SYR
and UNC as fully separate rows/sections, since a coach needs opponent-
specific numbers, not a blended average across both teams.

Sheets, in reading order:
  1. Game Averages        - per-team totals/averages, result breakdown
  2. Quarter Breakdown     - per-quarter, per-team possession counts/duration
  3. Cuts Summary          - per-team cut totals, plus per-quarter detail
  4. Formations by Result  - per-TEAM, per-result offense formation breakdown
  5. Cuts by Result        - per-TEAM, per-result average cuts
  6. Possession Log        - full detail, one row per possession (read last)

Usage:
    python export_to_excel.py detected_possessions_final.json \\
        possession_results.json formation_analysis.json cuts_analysis.json

Output:
    game_report.xlsx
"""
import sys
import json
from pathlib import Path
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pipeline_common as pc

pc.pin_working_directory()

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TEAM_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
BOLD = Font(bold=True)


def load_json(path):
    with open(path) as f:
        return json.load(f)


def key_of(d):
    return (d['quarter'], d['team'], d['start_sec'])


def style_header_row(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def style_team_label(cell):
    cell.font = BOLD
    cell.fill = TEAM_FILL


def autofit_columns(ws, min_width=8, max_width=45):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 2, min_width), max_width)


def build_combined_rows(possessions, results_by_key, formations_by_key, cuts_by_key):
    rows = []
    for p in possessions:
        k = key_of(p)
        result_entry = results_by_key.get(k, {})
        formation_entry = formations_by_key.get(k, {})
        cuts_entry = cuts_by_key.get(k, {})

        rows.append({
            'Quarter': p['quarter'],
            'Team': p['team'],
            'Camera': p['camera'],
            'Start (s)': p['start_sec'],
            'End (s)': p['end_sec'],
            'Duration (s)': p['duration_sec'],
            'Result': result_entry.get('result', 'Untagged'),
            'Man Situation': result_entry.get('man_situation', ''),
            'Notes': result_entry.get('notes', ''),
            'Offense Formation': str(formation_entry.get('offense_dominant_formation', '')),
            'Defense Formation': str(formation_entry.get('defense_dominant_formation', '')),
            'Cuts In': cuts_entry.get('n_cuts_in', 0),
            'Cuts Out': cuts_entry.get('n_cuts_out', 0),
        })
    rows.sort(key=lambda r: (r['Team'], r['Quarter'], r['Start (s)']))
    return rows


def teams_in(rows):
    return sorted(set(r['Team'] for r in rows))


def write_game_averages(wb, rows):
    ws = wb.active
    ws.title = "Game Averages"
    ws.append(["GAME AVERAGES - per team (never combined)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])

    for team in teams_in(rows):
        team_rows = [r for r in rows if r['Team'] == team]
        n = len(team_rows)
        total_dur = sum(r['Duration (s)'] for r in team_rows)
        total_cuts_in = sum(r['Cuts In'] for r in team_rows)
        total_cuts_out = sum(r['Cuts Out'] for r in team_rows)
        # Substring match ("Goal" appears in "Goal", "Goal - Penalty Shot", etc.)
        # so every goal-type result counts toward one combined total, rather
        # than needing the reader to manually add several separately-listed
        # result rows together (e.g. "Goal" and "Goal - Penalty Shot" were
        # previously two separate lines below, easy to undercount by reading
        # only one of them).
        total_goals = sum(1 for r in team_rows if 'goal' in r['Result'].lower())

        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=team)
        style_team_label(ws.cell(row=r, column=1))
        ws.append(["Total Possessions", n])
        ws.append(["Total Goals", total_goals])
        ws.append(["Total Possession Time (s)", round(total_dur, 1)])
        ws.append(["Avg Possession Length (s)", round(total_dur / n, 1) if n else 0])
        ws.append(["Total Cuts (In + Out)", total_cuts_in + total_cuts_out])
        ws.append(["Avg Cuts per Possession", round((total_cuts_in + total_cuts_out) / n, 1) if n else 0])
        ws.append([])

        result_counts = defaultdict(int)
        for rr in team_rows:
            if rr['Result'] != 'Untagged':
                result_counts[rr['Result']] += 1
        if result_counts:
            ws.append(["Result Breakdown", "Count", "%"])
            style_header_row(ws, row=ws.max_row)
            tagged_total = sum(result_counts.values())
            for result, count in sorted(result_counts.items(), key=lambda kv: -kv[1]):
                ws.append([result, count, round(count / tagged_total * 100, 1)])
        ws.append([])
        ws.append([])

    autofit_columns(ws)


def write_quarter_breakdown(wb, rows):
    ws = wb.create_sheet("Quarter Breakdown")
    ws.append(["Team", "Quarter", "# Possessions", "Total Duration (s)", "Avg Duration (s)"])
    style_header_row(ws)
    by_team_quarter = defaultdict(list)
    for r in rows:
        by_team_quarter[(r['Team'], r['Quarter'])].append(r)

    for team in teams_in(rows):
        for (t, q), qrows in sorted(by_team_quarter.items()):
            if t != team:
                continue
            n = len(qrows)
            total_dur = sum(r['Duration (s)'] for r in qrows)
            ws.append([team, q, n, round(total_dur, 1), round(total_dur / n, 1) if n else 0])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autofit_columns(ws)


def write_cuts_summary(wb, rows):
    ws = wb.create_sheet("Cuts Summary")
    ws.append(["CUTS SUMMARY - per team (never combined)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])

    for team in teams_in(rows):
        team_rows = [r for r in rows if r['Team'] == team]
        n = len(team_rows)
        total_in = sum(r['Cuts In'] for r in team_rows)
        total_out = sum(r['Cuts Out'] for r in team_rows)

        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=team)
        style_team_label(ws.cell(row=r, column=1))
        ws.append(["Total Cuts In (Perimeter -> Net-to-8)", total_in])
        ws.append(["Total Cuts Out (Net-to-8 -> Perimeter)", total_out])
        ws.append(["Total Cuts", total_in + total_out])
        ws.append(["Avg Cuts per Possession", round((total_in + total_out) / n, 1) if n else 0])
        ws.append([])

        ws.append(["Quarter", "Cuts In", "Cuts Out", "Total Cuts"])
        style_header_row(ws, row=ws.max_row)
        by_q = defaultdict(lambda: [0, 0])
        for rr in team_rows:
            by_q[rr['Quarter']][0] += rr['Cuts In']
            by_q[rr['Quarter']][1] += rr['Cuts Out']
        for q, (ci, co) in sorted(by_q.items()):
            ws.append([q, ci, co, ci + co])
        ws.append([])
        ws.append([])

    autofit_columns(ws)


def write_formations_by_result(wb, rows):
    ws = wb.create_sheet("Formations by Result")
    ws.append(["Team", "Result", "Offense Formation (Behind Net, Net-to-8, Perimeter)", "Count", "% within Result"])
    style_header_row(ws)

    by_team_result = defaultdict(lambda: defaultdict(int))
    for r in rows:
        if r['Result'] == 'Untagged' or not r['Offense Formation']:
            continue
        by_team_result[(r['Team'], r['Result'])][r['Offense Formation']] += 1

    for team in teams_in(rows):
        for (t, result), formations in sorted(by_team_result.items()):
            if t != team:
                continue
            total = sum(formations.values())
            for formation, count in sorted(formations.items(), key=lambda kv: -kv[1]):
                ws.append([team, result, formation, count, round(count / total * 100, 1)])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autofit_columns(ws)


def write_cuts_by_result(wb, rows):
    ws = wb.create_sheet("Cuts by Result")
    ws.append(["Team", "Result", "# Possessions", "Avg Cuts In", "Avg Cuts Out",
               "Avg Total Cuts", "Avg Duration (s)"])
    style_header_row(ws)

    by_team_result = defaultdict(lambda: {'in': [], 'out': [], 'dur': []})
    for r in rows:
        if r['Result'] == 'Untagged':
            continue
        d = by_team_result[(r['Team'], r['Result'])]
        d['in'].append(r['Cuts In'])
        d['out'].append(r['Cuts Out'])
        d['dur'].append(r['Duration (s)'])

    for team in teams_in(rows):
        for (t, result), d in sorted(by_team_result.items(), key=lambda kv: (-len(kv[1]['in']))):
            if t != team:
                continue
            n = len(d['in'])
            avg_in = sum(d['in']) / n
            avg_out = sum(d['out']) / n
            avg_dur = sum(d['dur']) / n
            ws.append([team, result, n, round(avg_in, 1), round(avg_out, 1),
                       round(avg_in + avg_out, 1), round(avg_dur, 1)])
    ws.freeze_panes = "A2"
    autofit_columns(ws)


def write_possession_log(wb, rows):
    ws = wb.create_sheet("Possession Log")
    if not rows:
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for r in rows:
        ws.append([r[h] for h in headers])
    style_header_row(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autofit_columns(ws)


def main():
    if len(sys.argv) < 5:
        print("Usage: python export_to_excel.py detected_possessions_final.json "
              "possession_results.json formation_analysis.json cuts_analysis.json")
        sys.exit(1)

    possessions = load_json(sys.argv[1])
    results = load_json(sys.argv[2]) if Path(sys.argv[2]).exists() else []
    formations = load_json(sys.argv[3]) if Path(sys.argv[3]).exists() else []
    cuts = load_json(sys.argv[4]) if Path(sys.argv[4]).exists() else []

    results_by_key = {key_of(r): r for r in results}
    formations_by_key = {key_of(f): f for f in formations}
    cuts_by_key = {key_of(c): c for c in cuts}

    rows = build_combined_rows(possessions, results_by_key, formations_by_key, cuts_by_key)

    untagged = sum(1 for r in rows if r['Result'] == 'Untagged')
    if untagged:
        print(f"Note: {untagged}/{len(rows)} possessions are untagged (no result) - "
              f"they'll show 'Untagged' in the log and be excluded from the "
              f"result-based breakdown sheets.")

    wb = Workbook()
    # Ordered simple -> complex for a coach reading top to bottom:
    write_game_averages(wb, rows)
    write_quarter_breakdown(wb, rows)
    write_cuts_summary(wb, rows)
    write_formations_by_result(wb, rows)
    write_cuts_by_result(wb, rows)
    write_possession_log(wb, rows)

    out_path = Path('game_report.xlsx')
    wb.save(out_path)
    print(f"\n\u2713 Saved {out_path} ({len(rows)} possessions across 6 sheets)")


if __name__ == '__main__':
    main()
